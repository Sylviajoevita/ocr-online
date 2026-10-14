from flask import Flask, request, send_file, render_template_string
from flask_cors import CORS
from PIL import Image
import pytesseract, io, json, os, traceback
from openpyxl import Workbook
from openpyxl.styles import Font

app = Flask(__name__)
CORS(app)

HTML_PAGE = """<!DOCTYPE html><html><head><meta charset="UTF-8"><meta name="viewport" content="width=device-width, initial-scale=1.0"><title>OCR Studio</title>
<style>*{box-sizing:border-box;font-family:sans-serif}body{margin:0;min-height:100vh;display:flex;align-items:center;justify-content:center;padding:20px;background:#7a8cff}.card{width:100%;max-width:700px;background:rgba(255,255,255,.3);backdrop-filter:blur(15px);border-radius:20px;padding:30px}.drop{border:1.5px dashed #fff;border-radius:14px;padding:30px;text-align:center;background:rgba(255,255,255,.4);cursor:pointer;font-weight:600}.preview{display:grid;grid-template-columns:repeat(auto-fill,minmax(100px,1fr));gap:10px;margin-top:15px}.item{background:#fff;border-radius:8px;padding:6px;text-align:center;font-size:10px;position:relative}.item img{width:100%;height:60px;object-fit:cover}.item button{position:absolute;top:2px;right:2px;background:#000;color:#fff;border:none;border-radius:50%;width:18px;height:18px}.btn{width:100%;padding:14px;background:#111;color:#fff;border-radius:10px;margin-top:15px;display:none;font-weight:700;border:none}.btn.show{display:block}.log{background:#111;color:#0f0;padding:10px;border-radius:8px;margin-top:10px;font-family:monospace;font-size:11px;white-space:pre-wrap}</style></head>
<body><div class="card"><h2 style="text-align:center;margin:0 0 10px">OCR Studio - Fast</h2>
<div class="drop" onclick="document.getElementById('files').click()">Click to select PDF, JSON, Images<input type="file" id="files" multiple hidden accept=".pdf,.json,image/*"></div>
<div id="preview" class="preview"></div><button class="btn" id="btn" onclick="upload()">Process & Download Excel</button><div id="log" class="log">Ready</div></div>
<script>let selectedFiles=[];const input=document.getElementById('files'),preview=document.getElementById('preview'),btn=document.getElementById('btn'),log=document.getElementById('log');
input.addEventListener('change',e=>{selectedFiles=selectedFiles.concat(Array.from(e.target.files));render();});
function render(){preview.innerHTML='';if(!selectedFiles.length){btn.classList.remove('show');return;}selectedFiles.forEach((file,i)=>{const d=document.createElement('div');d.className='item';d.innerHTML=file.type.startsWith('image/')?`<img src="${URL.createObjectURL(file)}">`:(file.name.endsWith('.pdf')?'📕 PDF':'📄 JSON');const n=document.createElement('div');n.textContent=file.name;const rm=document.createElement('button');rm.textContent='×';rm.onclick=()=>{selectedFiles.splice(i,1);render();};d.appendChild(n);d.appendChild(rm);preview.appendChild(d);});btn.classList.add('show');btn.textContent=`Process ${selectedFiles.length} file(s)`;}
async function upload(){if(!selectedFiles.length)return;btn.disabled=true;btn.textContent='Processing...';log.textContent='Processing...';const fd=new FormData();selectedFiles.forEach(f=>fd.append('files',f));try{const res=await fetch('/ocr',{method:'POST',body:fd});if(!res.ok){const t=await res.text();log.textContent=t.slice(0,3000);throw new Error('Server error');}const blob=await res.blob();const url=URL.createObjectURL(blob);const a=document.createElement('a');a.href=url;a.download='extracted.xlsx';a.click();log.textContent='✅ Done';btn.textContent='✅ Done';}catch(e){log.textContent='Error: '+e.message;btn.textContent='Try Again';}finally{btn.disabled=false;}}</script></body></html>
"""

@app.route('/')
def home(): return render_template_string(HTML_PAGE)
@app.route('/background.jpg')
def bg(): return send_file("background.jpg", mimetype="image/jpeg") if os.path.exists("background.jpg") else ("",404)
@app.route('/favicon.ico')
def fav(): return send_file("favicon.ico", mimetype="image/x-icon") if os.path.exists("favicon.ico") else ("",404)

def fast_ocr(img):
    if max(img.size) > 1500: img.thumbnail((1500,1500))
    gray = img.convert("L")
    try:
        d = pytesseract.image_to_data(gray, lang='eng', config='--oem 1 --psm 6', output_type=pytesseract.Output.DICT)
        lines = {}
        for i in range(len(d['text'])):
            t = d['text'][i].strip()
            if not t: continue
            try:
                if int(d['conf'][i]) < 35: continue
            except: continue
            k = d['line_num'][i]
            if k not in lines: lines[k] = []
            lines[k].append({'text': t, 'left': d['left'][i], 'width': d['width'][i]})
        rows = []
        for k in sorted(lines):
            words = sorted(lines[k], key=lambda w: w['left'])
            avg = sum(w['width'] for w in words)/len(words)
            cols = []; cur = words[0]['text']; last = words[0]['left']+words[0]['width']
            for w in words[1:]:
                gap = w['left']-last
                if gap > avg: cols.append(cur); cur = w['text']
                else: cur += " "+w['text']
                last = w['left']+w['width']
            cols.append(cur); rows.append(cols)
        if rows: return rows
    except: pass
    txt = pytesseract.image_to_string(gray, lang='eng', config='--oem 1 --psm 6')
    return [[l] for l in txt.splitlines() if l.strip()]

@app.route('/ocr', methods=['POST'])
def ocr():
    try:
        files = request.files.getlist('files')
        if not files: return "No files", 400
        wb = Workbook(); wb.remove(wb.active)
        for f in files:
            name = f.filename or "file"
            ext = name.rsplit('.',1)[-1].lower() if '.' in name else ""
            safe = "".join(c if c.isalnum() or c in (" ","_") else "_" for c in name.rsplit('.',1)[0])[:31] or "Sheet1"
            ws = wb.create_sheet(title=safe)
            ws['A1'] = name; ws['A1'].font = Font(bold=True, color="0000FF")
            try:
                if ext == 'json':
                    data = json.loads(f.read().decode('utf-8'))
                    if isinstance(data, list) and data and isinstance(data[0], dict):
                        headers = list(data[0].keys())
                        for ci,h in enumerate(headers, start=1): ws.cell(row=3, column=ci, value=h).font = Font(bold=True)
                        for ri, rd in enumerate(data, start=4):
                            for ci,h in enumerate(headers, start=1): ws.cell(row=ri, column=ci, value=str(rd.get(h,"")))
                    else:
                        ws.cell(row=3, column=1, value=str(data))
                elif ext == 'pdf':
                    import PyPDF2
                    reader = PyPDF2.PdfReader(io.BytesIO(f.read()))
                    r = 3
                    for pi in range(min(len(reader.pages), 15)):
                        page = reader.pages[pi]
                        text = page.extract_text() or ""
                        ws.cell(row=r, column=1, value=f"Page {pi+1}").font = Font(bold=True); r+=1
                        for line in text.splitlines():
                            if not line.strip(): continue
                            parts = line.split()
                            if len(parts)>1:
                                for ci,p in enumerate(parts, start=1): ws.cell(row=r, column=ci, value=p)
                            else:
                                ws.cell(row=r, column=1, value=line.strip())
                            r+=1
                        r+=1
                else:
                    img = Image.open(f.stream)
                    rows = fast_ocr(img)
                    r = 3
                    for row_data in rows:
                        for ci,v in enumerate(row_data, start=1): ws.cell(row=r, column=ci, value=v)
                        r+=1
                for c in ['A','B','C','D','E','F']: ws.column_dimensions[c].width=25
            except Exception as inner:
                ws.cell(row=3, column=1, value=f"Error: {inner}")
        out = io.BytesIO(); wb.save(out); out.seek(0)
        return send_file(out, as_attachment=True, download_name="extracted.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    except Exception as e:
        return f"Server Error: {e}\n{traceback.format_exc()}", 500

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=int(os.environ.get("PORT", 10000)))
